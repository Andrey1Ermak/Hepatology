#!/usr/bin/env python3
"""Daily sync with the lesson repository (Andrey1Ermak/reading-final).

Usage (from the Hepatology root):
  python3 tools/sync_sources.py              # pull, report and extract new/changed lessons
  python3 tools/sync_sources.py --done       # record the current lesson commit as processed
  python3 tools/sync_sources.py --extract 02-03 02-07   # extract specific lessons by prefix
  python3 tools/sync_sources.py --extract 02-            # every lesson of block 2

Lessons live in sources/reading-final/hepatology/ as BB-DD_Topic.pptx (deck) and
BB-DD_Topic_audio.pdf (chapter text). Extracted text goes to sources/extracted/
(regenerable, not committed). Processing state lives in sources/sync_state.json.
"""
import json
import subprocess
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
REPO_URL = "https://github.com/Andrey1Ermak/reading-final.git"
CLONE = ROOT / "sources" / "reading-final"
LESSONS = "hepatology"
PLAN = "Hepatology_Year_Plan.docx"
STATE = ROOT / "sources" / "sync_state.json"
EXTRACTED = ROOT / "sources" / "extracted"


def git(*args):
    return subprocess.run(["git", "-C", str(CLONE), *args], check=True, capture_output=True, text=True).stdout.strip()


def load_state():
    return json.loads(STATE.read_text()) if STATE.exists() else {"last_commit": None, "last_sync": None}


def pull():
    if not CLONE.exists():
        subprocess.run(["git", "clone", "-q", REPO_URL, str(CLONE)], check=True)
    else:
        git("pull", "-q", "--ff-only")
    return git("rev-parse", "HEAD")


def changed_since(commit):
    """(status, path) for lesson files and the year plan changed after `commit`."""
    paths = [LESSONS, PLAN]
    if commit is None:
        out = git("ls-files", *paths)
        return [("A", p) for p in out.splitlines() if p]
    out = git("diff", "--name-status", f"{commit}..HEAD", "--", *paths)
    return [tuple(line.split("\t", 1)) for line in out.splitlines() if line]


def extract_pptx(path):
    from pptx import Presentation

    lines = [f"# {path.stem}", ""]
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        lines.append(f"## Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if text:
                    lines.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"> Notes: {notes}")
        lines.append("")
    return "\n".join(lines)


def extract_pdf(path):
    return subprocess.run(["pdftotext", "-layout", str(path), "-"], check=True, capture_output=True, text=True).stdout


def extract(rel_path):
    src = CLONE / rel_path
    EXTRACTED.mkdir(parents=True, exist_ok=True)
    if src.suffix == ".pptx":
        text = extract_pptx(src)
    elif src.suffix == ".pdf":
        text = extract_pdf(src)
    else:
        return None
    out = EXTRACTED / f"{src.stem}.md"
    out.write_text(text)
    return out.relative_to(ROOT)


def lesson_id(rel_path):
    """'hepatology/02-07_Metabolic_citrin.pptx' -> (2, 7, 'Metabolic citrin')."""
    stem = Path(rel_path).stem.removesuffix("_audio")
    prefix, _, topic = stem.partition("_")
    block, _, day = prefix.partition("-")
    return int(block), int(day), topic.replace("_", " ")


def main(argv):
    state = load_state()

    if "--done" in argv:
        state["last_commit"] = git("rev-parse", "HEAD")
        state["last_sync"] = date.today().isoformat()
        STATE.write_text(json.dumps(state, indent=2) + "\n")
        print(f"Recorded lessons commit {state['last_commit'][:7]} as processed ({state['last_sync']}).")
        return

    if "--extract" in argv:
        prefixes = argv[argv.index("--extract") + 1:]
        files = sorted(p for p in (CLONE / LESSONS).iterdir() if any(p.name.startswith(x) for x in prefixes))
        for f in files:
            print(extract(f.relative_to(CLONE)))
        if not files:
            print("No lessons match", prefixes)
        return

    head = pull()
    changes = changed_since(state["last_commit"])
    print(f"Lessons repo HEAD: {head[:7]} · last processed: {(state['last_commit'] or 'never')[:7]} on {state['last_sync']}")
    if not changes:
        print("NO NEW MATERIAL — nothing to sync today.")
        return

    plan_changed = any(p == PLAN for _, p in changes)
    lessons = {}
    for status, path in changes:
        if path.startswith(LESSONS + "/") and status != "D":
            lessons.setdefault(lesson_id(path), []).append(extract(path))
    for (block, day, topic), outs in sorted(lessons.items()):
        print(f"NEW/CHANGED  block {block} day {day}: {topic}")
        for out in outs:
            print(f"    {out}")
    for status, path in changes:
        if status == "D":
            print(f"DELETED  {path}")
    if plan_changed:
        print(f"YEAR PLAN CHANGED: {PLAN} — update Content/Registry/plan.json (if it exists yet).")
    print("After integrating, run: python3 tools/sync_sources.py --done")


if __name__ == "__main__":
    main(sys.argv[1:])
